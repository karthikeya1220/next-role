"""Turn an uploaded career document into proposed KB chunks.

Two clearly separated steps:
1. extract_text  — deterministic: pull raw text out of pdf/docx/txt/md.
2. parse_to_chunks — the LLM's job: structure that text into accomplishment
   chunks. It is instructed to use ONLY what's in the document (no inventing);
   the user reviews and edits everything before it is saved.
"""
import io
from typing import Callable

import pdfplumber
from docx import Document

from app.ai.llm import generate_json

# How much document text goes into one model call. A long CV is split across
# several calls rather than truncated: cutting at 12k characters silently threw
# away the second half of anyone's career.
_SEGMENT_CHARS = 12000


def extract_text(filename: str, data: bytes) -> str:
    """Extract plain text from a supported document."""
    ext = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    if ext == "pdf":
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            return "\n".join(page.extract_text() or "" for page in pdf.pages)
    if ext == "docx":
        doc = Document(io.BytesIO(data))
        return "\n".join(p.text for p in doc.paragraphs)
    if ext in ("txt", "md"):
        return data.decode("utf-8", errors="ignore")
    raise ValueError(f"Unsupported file type: .{ext} (use pdf, docx, txt, or md)")


_SYSTEM = """You extract a candidate's career into structured accomplishment chunks for a resume knowledge base.

Rules:
- Use ONLY information present in the document. Never invent facts, metrics, employers, dates, or technologies.
- Create ONE chunk per distinct accomplishment or responsibility (split multi-part bullets).
- Choose `type` from: project, experience, leadership, achievement, skill, certification.
- Fill `technologies` and `skills` only with items actually mentioned; otherwise use [].
- Fill `impact` only if a concrete outcome/metric is stated; otherwise null.
- Keep `accomplishment` concise, truthful, and results-oriented.

Respond as JSON: {"chunks": [{"type","title","context","company","date_range","accomplishment","technologies":[],"skills":[],"impact"}]}"""


def segment(text: str) -> list[str]:
    """Split a document into model-sized pieces, breaking at the cleanest seam.

    Where it breaks matters: cutting mid-bullet hands the model half an
    accomplishment and it faithfully records half an accomplishment. So it tries
    paragraph gaps first, then single line breaks, and only chops mid-line when a
    document offers no seam at all.

    Both fallbacks are load-bearing on real files. A Windows text file separates
    paragraphs with CRLF, and PDF text extraction routinely comes back as one
    long run of single newlines — either would otherwise sail past a
    blank-line-only split and reach the model whole.
    """
    text = text.replace("\r\n", "\n").replace("\r", "\n").strip()
    if len(text) <= _SEGMENT_CHARS:
        return [text] if text else []

    return _pack(_pieces(text))


def _pieces(text: str) -> list[str]:
    """Break text into units no larger than one segment, cleanest seam first."""
    units = text.split("\n\n")
    if all(len(u) <= _SEGMENT_CHARS for u in units):
        return units

    units = [line for unit in units for line in unit.split("\n")]
    if all(len(u) <= _SEGMENT_CHARS for u in units):
        return units

    return [
        unit[i : i + _SEGMENT_CHARS]
        for unit in units
        for i in range(0, len(unit), _SEGMENT_CHARS)
    ]


def _pack(units: list[str]) -> list[str]:
    """Greedily fill segments so a long document is as few model calls as possible."""
    segments: list[str] = []
    current = ""
    for unit in units:
        if not unit.strip():
            continue
        if current and len(current) + len(unit) + 2 > _SEGMENT_CHARS:
            segments.append(current)
            current = unit
        else:
            current = f"{current}\n\n{unit}" if current else unit
    if current:
        segments.append(current)
    return segments


def parse_to_chunks(
    text: str,
    kind_hint: str | None = None,
    on_segment: Callable[[], None] | None = None,
) -> list[dict]:
    """Structure document text into chunk dicts, one model call per segment.

    No timeout: a 30-page CV on a laptop model is slow, not broken, and the
    caller runs this in the background with a progress bar for exactly that
    reason. `on_segment` fires after each call so that bar can move.
    """
    hint = f"\n\nThis document is the candidate's: {kind_hint}." if kind_hint else ""
    chunks: list[dict] = []
    for piece in segment(text):
        # A resume can yield many chunks, so this needs far more room than extraction.
        result = generate_json(
            _SYSTEM, f"Document:{hint}\n\n{piece}", timeout=None, max_tokens=-1
        )
        raw = result.get("chunks", []) if isinstance(result, dict) else []
        chunks.extend(
            _normalize(c) for c in raw if isinstance(c, dict) and c.get("accomplishment")
        )
        if on_segment:
            on_segment()
    return chunks


_VALID_TYPES = {"project", "experience", "leadership", "achievement", "skill", "certification"}

# LLMs often emit the *string* "null"/"none"/"n/a" instead of a real null.
_NULLISH = {"", "null", "none", "n/a", "na", "unknown", "not specified"}


def _opt(value, limit: int) -> str | None:
    """Optional text field: strip, drop nullish placeholders, cap length."""
    if value is None:
        return None
    text = str(value).strip()
    return text[:limit] if text.lower() not in _NULLISH else None


def _normalize(c: dict) -> dict:
    """Coerce a raw LLM chunk into our exact shape with safe defaults."""
    ctype = str(c.get("type", "")).lower().strip()
    return {
        "type": ctype if ctype in _VALID_TYPES else "experience",
        "title": _opt(c.get("title"), 300) or "Untitled",
        "context": _opt(c.get("context"), 300),
        "company": _opt(c.get("company"), 200),
        "date_range": _opt(c.get("date_range"), 100),
        "accomplishment": str(c["accomplishment"]).strip(),
        "technologies": [t for t in (_opt(x, 100) for x in (c.get("technologies") or [])) if t],
        "skills": [s for s in (_opt(x, 100) for x in (c.get("skills") or [])) if s],
        "impact": _opt(c.get("impact"), 500),
    }
